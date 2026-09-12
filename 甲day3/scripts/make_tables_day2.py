"""
题设表 3 / 表 4 / 表 5 生成器                 [甲 · M0 · Day2]

为什么用脚本而不是手抄
----------------------
分工文档明确要求："**禁止只手改论文中的干燥时间**；
结果变更时：先更新配置标识和结果 → 自动重建表图 → 最后改论文。"
本脚本从**已写出的 result2.xlsx / result3.xlsx** 反读并抽表，
因此论文里的表与提交的结果文件**必然一致**。

题面要求（[Q-PDF] 逐字）
------------------------
表3/表4：3 h 内每隔 0.5 h、到药材中心距离 0、0.5、1、1.5、2 cm  → 6×5
表5    ：每隔 6 h、到药材中心距离每隔 0.5 cm；**末行为"烘干结束时间"**

⚠️ 表5 的末行**不是** 6 h 的整数倍（连续模型的达标时刻一般不落在整点上），
   题面用"烘干结束时间"这一行正是为它留的。
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.config import DOC_DIR, R_OUT_CM  # noqa: E402

COLS_CM = [0.0, 0.5, 1.0, 1.5, 2.0]


def _read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    t = np.array([wb[wb.sheetnames[0]].cell(2 + i, 1).value
                  for i in range(wb[wb.sheetnames[0]].max_row - 1)], dtype=float)
    out = {}
    for s in wb.sheetnames:
        # ⚠️ 必须按名字白名单筛选：`运行信息` sheet 也是 2 列，
        #    但第 2 列存的是文本（如 "M0"），按列数过滤会把它读进来并炸掉 float()。
        if s not in ("温度", "水分浓度", "Sheet1"):
            continue
        ws = wb[s]
        if ws.max_row < 2 or ws.max_column < len(R_OUT_CM) + 1:
            continue
        arr = np.array([[ws.cell(2 + i, 2 + j).value
                         for j in range(len(R_OUT_CM))]
                        for i in range(ws.max_row - 1)], dtype=float)
        if arr.shape[0] == len(t):
            out[s] = arr
    return t, out


def _pick_cols(arr):
    """按 0/0.5/1/1.5/2 cm 取列。用 ... 索引以便同时接受 1-D 行与 2-D 矩阵。"""
    idx = [R_OUT_CM.index(c) for c in COLS_CM]
    return np.asarray(arr)[..., idx]


def _row_at(t, arr, t_target, tol=0.51):
    k = int(np.argmin(np.abs(t - t_target)))
    assert abs(t[k] - t_target) <= tol, \
        f"结果文件里没有 t={t_target} 的行（最近的 {t[k]}）"
    return k


def _md_table(header, rows, first_col_name):
    out = ["| " + first_col_name + " | " + " | ".join(header) + " |",
           "|" + "---|" * (len(header) + 1)]
    for lab, vals in rows:
        out.append("| " + lab + " | " + " | ".join(vals) + " |")
    return "\n".join(out)


def build(layer="M0"):
    r2 = ROOT / "results" / layer / "result2.xlsx"
    r3 = ROOT / "results" / layer / "result3.xlsx"
    if not r2.exists() or not r3.exists():
        raise FileNotFoundError(f"缺 {r2} 或 {r3}")

    t2, d2 = _read(r2)
    t3, d3 = _read(r3)
    hdr = [f"{c:g}" for c in COLS_CM]
    lines = [f"# 题设表 3 / 表 4 / 表 5（{layer}）", "",
             "> 由 `scripts/make_tables_day2.py` 从 result2.xlsx / result3.xlsx "
             "**反读生成**，与提交文件逐格一致。",
             "> 全部数值保留 4 位小数（题面要求）。", ""]

    # ---- 表3 温度 ----
    lines += ["## 表 3　3 小时内药材的温度（单位：°C）", ""]
    rows = []
    for hh in (0.5, 1.0, 1.5, 2.0, 2.5, 3.0):
        k = _row_at(t2, d2["温度"], hh * 3600.0)
        rows.append((f"{hh:.1f}", [f"{v:.4f}" for v in _pick_cols(d2['温度'][k])]))
    lines += [_md_table(hdr, rows, "时间/h"), ""]

    # ---- 表4 水分浓度 ----
    lines += ["## 表 4　3 小时内药材的水分浓度（单位：kg/kg）", ""]
    rows = []
    for hh in (0.5, 1.0, 1.5, 2.0, 2.5, 3.0):
        k = _row_at(t2, d2["水分浓度"], hh * 3600.0)
        rows.append((f"{hh:.1f}", [f"{v:.4f}" for v in _pick_cols(d2['水分浓度'][k])]))
    lines += [_md_table(hdr, rows, "时间/h"), ""]

    # ---- 表5 烘干过程水分浓度 + 末行达标时刻 ----
    C3 = d3["Sheet1"] if "Sheet1" in d3 else list(d3.values())[0]
    t_end_s = float(t3[-1])
    lines += ["## 表 5　药材烘干过程的水分浓度（单位：kg/kg）", ""]
    rows = []
    hh = 6.0
    while hh * 3600.0 <= t_end_s + 1e-9:
        k = _row_at(t3, C3, hh * 3600.0, tol=0.51)
        rows.append((f"{hh:.0f}", [f"{v:.4f}" for v in _pick_cols(C3[k])]))
        hh += 6.0
    kl = len(t3) - 1
    rows.append(("**烘干结束时间**",
                 [f"{v:.4f}" for v in _pick_cols(C3[kl])]))
    lines += [_md_table(hdr, rows, "时间/h"), ""]
    lines += [f"> **烘干结束时间 = {t_end_s:.0f} s = {t_end_s/3600:.4f} h "
              f"= {t_end_s/86400:.4f} 天**（result3.xlsx 的末行；"
              f"未舍入值见 `docs/day2_core.json` 的 `t_star_s`）。",
              "> 该行**不落在 6 h 的整数倍上**，正是题面为它单列一行的原因。", ""]
    return "\n".join(lines), t_end_s


def main():
    layer = sys.argv[1] if len(sys.argv) > 1 else "M0"
    txt, t_end = build(layer)
    p = DOC_DIR / f"Day2_题设表3-4-5_{layer}.md"
    p.write_text(txt, encoding="utf-8")
    print(f"  → {p}   (t* = {t_end:.0f} s = {t_end/3600:.4f} h)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
