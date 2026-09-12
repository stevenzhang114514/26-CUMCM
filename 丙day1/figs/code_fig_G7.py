"""
G7 问题4：物理坐标下的含水率演化 + 移动边界轨迹
输入：
    results/M0/result4.xlsx
    results/raw/radius_fit_p4_v1.csv（等乙提供）
输出：
    figs/G7_p4_moving_boundary.png
    figs/G7_p4_moving_boundary.pdf
"""

from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parent.parent
M0_DIR = ROOT / "results" / "M0"
RAW_DIR = ROOT / "results" / "raw"
FIGS_DIR = ROOT / "figs"
FIGS_DIR.mkdir(parents=True, exist_ok=True)

plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

SHOW_TIMES_H = [0, 6, 12, 24, 48, 72]


def load_sheet(xlsx_path, sheet_name):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    header = [ws.cell(row=1, column=c).value
              for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value
               for c in range(1, ws.max_column + 1)]
        rows.append(row)
    wb.close()
    return header, rows


def load_fit():
    path = RAW_DIR / "radius_fit_p4_v1.csv"
    if not path.exists():
        print(f"[缺失] {path}，等乙提供")
        return None
    return pd.read_csv(path)


def main():
    result4 = M0_DIR / "result4.xlsx"
    if not result4.exists():
        print(f"[缺失] {result4}")
        return

    header, rows = load_sheet(result4, "Sheet1")

    # A 列时间
    times = np.array([r[0] for r in rows], dtype=float)
    # 距离列：能转 float 的
    r_cols = []
    surface_col = None
    for j, c in enumerate(header):
        if j == 0:
            continue
        try:
            float(c)
            r_cols.append((j, float(c)))
        except (ValueError, TypeError):
            if str(c).strip() == "药材表面":
                surface_col = j

    r_cols.sort(key=lambda x: x[1])

    fit = load_fit()

    fig, axes = plt.subplots(1, 2, figsize=(14, 5))

    # 左图：不同时刻的含水率径向分布
    for th in SHOW_TIMES_H:
        tt = th * 3600
        idx = int(np.argmin(np.abs(times - tt)))
        rr = []
        cc = []
        for j, rv in r_cols:
            v = rows[idx][j]
            if v is None:
                continue
            try:
                v = float(v)
                if np.isnan(v):
                    continue
                rr.append(rv)
                cc.append(v)
            except (ValueError, TypeError):
                continue
        axes[0].plot(rr, cc, "-o", ms=3, lw=1.2, label=f"t = {th} h")

    axes[0].set_xlabel("到药材中心的距离 / cm")
    axes[0].set_ylabel("水分浓度 / (kg/kg)")
    axes[0].set_title("(a) 物理坐标下的含水率演化")
    axes[0].legend()
    axes[0].grid(alpha=0.3)

    # 右图：移动边界 R(t)
    if fit is not None:
        axes[1].plot(fit["time_s"], fit["radius_cm"], "-", lw=1.5,
                     label="R(t)")
    else:
        axes[1].text(0.5, 0.5, "等乙提供 radius_fit_p4_v1.csv",
                     ha="center", va="center",
                     transform=axes[1].transAxes)
    axes[1].set_xlabel("时间 / s")
    axes[1].set_ylabel("半径 / cm")
    axes[1].set_title("(b) 移动边界轨迹")
    axes[1].legend()
    axes[1].grid(alpha=0.3)

    fig.suptitle("G7 问题4 物理坐标与移动边界")
    fig.tight_layout()

    out_png = FIGS_DIR / "G7_p4_moving_boundary.png"
    out_pdf = FIGS_DIR / "G7_p4_moving_boundary.pdf"
    fig.savefig(out_png, dpi=200)
    fig.savefig(out_pdf)
    plt.close(fig)
    print("G7 已保存:", out_png)
    print("G7 已保存:", out_pdf)


if __name__ == "__main__":
    main()