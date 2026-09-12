"""
G6 问题3：C_max(t) 衰减曲线 + 0.15 阈值线 + 达标点 + 终态径向分布
直接读 results/M0/result3.xlsx。
输出：
    figs/G6_p3_threshold.png
    figs/G6_p3_threshold.pdf
"""

from pathlib import Path
import numpy as np
import openpyxl
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parent.parent
M0_DIR = ROOT / "results" / "M0"
FIGS_DIR = ROOT / "figs"
FIGS_DIR.mkdir(parents=True, exist_ok=True)

plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

THRESHOLD = 0.15


def load_sheet(xlsx_path, sheet_name):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    header = [ws.cell(row=1, column=c).value
              for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value
               for c in range(1, ws.max_column + 1)]
        rows.append(row)
    wb.close()
    data = np.array(rows, dtype=float)
    return header, data


def main():
    result3 = M0_DIR / "result3.xlsx"
    if not result3.exists():
        print(f"[缺失] {result3}")
        return

    header, data = load_sheet(result3, "Sheet1")
    t = data[:, 0]
    r_cols = [c for c in header if isinstance(c, (int, float))]
    r_vals = np.array([float(c) for c in r_cols])

    C = data[:, 1:]
    C_max = C.max(axis=1)

    idx_pass = np.where(C_max <= THRESHOLD)[0]
    if len(idx_pass) > 0:
        t_pass = t[idx_pass[0]]
        C_pass = C_max[idx_pass[0]]
    else:
        t_pass = None
        C_pass = None

    fig, axes = plt.subplots(1, 2, figsize=(14, 5))

    # 左：C_max(t) + 阈值
    axes[0].plot(t, C_max, "-", lw=1.5, label="$C_{max}(t)$")
    axes[0].axhline(THRESHOLD, color="red", ls="--",
                    label=f"阈值 = {THRESHOLD}")
    if t_pass is not None:
        axes[0].plot(t_pass, C_pass, "ro", ms=8,
                     label=f"达标 t = {t_pass:.0f} s")
    axes[0].set_xlabel("时间 / s")
    axes[0].set_ylabel("水分浓度 / (kg/kg)")
    axes[0].set_title("(a) 全域最大含水率与阈值")
    axes[0].legend()
    axes[0].grid(alpha=0.3)

    # 右：终态径向分布
    C_final = C[-1]
    axes[1].plot(r_vals, C_final, "-o", ms=3, lw=1.2)
    axes[1].axhline(THRESHOLD, color="red", ls="--")
    axes[1].set_xlabel("到药材中心的距离 / cm")
    axes[1].set_ylabel("水分浓度 / (kg/kg)")
    axes[1].set_title("(b) 终态径向分布")
    axes[1].grid(alpha=0.3)

    fig.suptitle("G6 问题3 阈值与终态")
    fig.tight_layout()

    out_png = FIGS_DIR / "G6_p3_threshold.png"
    out_pdf = FIGS_DIR / "G6_p3_threshold.pdf"
    fig.savefig(out_png, dpi=200)
    fig.savefig(out_pdf)
    plt.close(fig)
    print("G6 已保存:", out_png)
    print("G6 已保存:", out_pdf)

    if t_pass is not None:
        print(f"达标时间: {t_pass:.0f} s = {t_pass/3600:.2f} h")


if __name__ == "__main__":
    main()