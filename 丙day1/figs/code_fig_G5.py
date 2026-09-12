"""
G5 问题2 前3小时径向分布
直接读 results/M0/result2.xlsx。
输出：
    figs/G5_p2_radial.png
    figs/G5_p2_radial.pdf
"""

from pathlib import Path
import numpy as np
import openpyxl
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parent.parent
M0_DIR = ROOT / "results" / "M0"
FIGS_DIR = ROOT / "figs"
FIGS_DIR.mkdir(parents=True, exist_ok=True)

plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

# 展示时刻：0.5、1、2、3 h
SHOW_TIMES_H = [0.5, 1.0, 2.0, 3.0]


def load_sheet(xlsx_path, sheet_name):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    header = [ws.cell(row=1, column=c).value
              for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value
               for c in range(1, ws.max_column + 1)]
        rows.append(row)
    wb.close()
    data = np.array(rows, dtype=float)
    return header, data


def plot_radial(ax, header, data, ylabel, title):
    t = data[:, 0]
    r_cols = [c for c in header if isinstance(c, (int, float))]
    r_vals = np.array([float(c) for c in r_cols])

    for th in SHOW_TIMES_H:
        tt = th * 3600
        idx = np.argmin(np.abs(t - tt))
        y = data[idx, 1:]
        ax.plot(r_vals, y, "-o", ms=3, lw=1.2, label=f"t = {th} h")

    ax.set_xlabel("到药材中心的距离 / cm")
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.legend()
    ax.grid(alpha=0.3)


def main():
    result2 = M0_DIR / "result2.xlsx"
    if not result2.exists():
        print(f"[缺失] {result2}")
        return

    header_T, data_T = load_sheet(result2, "温度")
    header_C, data_C = load_sheet(result2, "水分浓度")

    fig, axes = plt.subplots(1, 2, figsize=(14, 5))

    plot_radial(axes[0], header_T, data_T, "温度 / ℃", "(a) 温度径向分布")
    plot_radial(axes[1], header_C, data_C, "水分浓度 / (kg/kg)", "(b) 含水率径向分布")

    fig.suptitle("G5 问题2 前3小时径向分布")
    fig.tight_layout()

    out_png = FIGS_DIR / "G5_p2_radial.png"
    out_pdf = FIGS_DIR / "G5_p2_radial.pdf"
    fig.savefig(out_png, dpi=200)
    fig.savefig(out_pdf)
    plt.close(fig)
    print("G5 已保存:", out_png)
    print("G5 已保存:", out_pdf)


if __name__ == "__main__":
    main()