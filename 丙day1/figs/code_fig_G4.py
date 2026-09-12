"""
G4 问题1 温度与含水率演化双图
分子图，禁用双纵轴。
直接读 results/M0/result1.xlsx。
输出：
    figs/G4_p1_evolution.png
    figs/G4_p1_evolution.pdf
"""

from pathlib import Path
import numpy as np
import openpyxl
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parent.parent
M0_DIR = ROOT / "results" / "M0"
FIGS_DIR = ROOT / "figs"
FIGS_DIR.mkdir(parents=True, exist_ok=True)

plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

# 展示用的位置：中心、中间、表面
SHOW_DIST = [0.0, 1.0, 2.0]


def load_sheet(xlsx_path, sheet_name):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]

    header = [ws.cell(row=1, column=c).value
              for c in range(1, ws.max_column + 1)]

    rows = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value
               for c in range(1, ws.max_column + 1)]
        rows.append(row)

    wb.close()

    data = np.array(rows, dtype=float)
    return header, data


def pick_col(header, target):
    cols = [c for c in header if isinstance(c, (int, float))]
    return min(cols, key=lambda c: abs(float(c) - target))


def plot_field(ax, header, data, ylabel, title):
    t = data[:, 0]
    for d in SHOW_DIST:
        col = pick_col(header, d)
        j = header.index(col)
        ax.plot(t, data[:, j], "-", lw=1.2,
                label=f"r = {float(col):.1f} cm")
    ax.set_xlabel("时间 / s")
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.legend()
    ax.grid(alpha=0.3)


def main():
    result1 = M0_DIR / "result1.xlsx"
    if not result1.exists():
        print(f"[缺失] {result1}")
        return

    header_T, data_T = load_sheet(result1, "温度")
    header_C, data_C = load_sheet(result1, "水分浓度")

    fig, axes = plt.subplots(2, 1, figsize=(10, 8), sharex=True)

    plot_field(axes[0], header_T, data_T, "温度 / ℃", "(a) 问题1 温度场演化")
    plot_field(axes[1], header_C, data_C, "水分浓度 / (kg/kg)", "(b) 问题1 含水率场演化")

    fig.suptitle("G4 问题1 温度与含水率演化")
    fig.tight_layout()

    out_png = FIGS_DIR / "G4_p1_evolution.png"
    out_pdf = FIGS_DIR / "G4_p1_evolution.pdf"
    fig.savefig(out_png, dpi=200)
    fig.savefig(out_pdf)
    plt.close(fig)
    print("G4 已保存:", out_png)
    print("G4 已保存:", out_pdf)


if __name__ == "__main__":
    main()