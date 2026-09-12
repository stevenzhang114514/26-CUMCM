"""
CODE-31b 生成完整 result1-4 模板
按题目要求生成完整时间轴、距离轴，覆盖原有残缺模板
生成到 data/ 下：
    result1_template.xlsx
    result2_template.xlsx
    result3_template.xlsx
    result4_template.xlsx
"""

from pathlib import Path
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)

# 距离轴：0, 0.1, 0.2, ..., 2.0
DIST = [round(i * 0.1, 1) for i in range(21)]


def write_sheet(ws, times, distances, last_col_name=None):
    """写入第1行距离、A列时间，其余留空。"""
    # 第1行
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, d in enumerate(distances):
        if last_col_name and j == len(distances) - 1:
            ws.cell(row=1, column=2 + j, value=last_col_name)
        else:
            ws.cell(row=1, column=2 + j, value=float(d))

    # A 列时间
    for i, t in enumerate(times):
        ws.cell(row=2 + i, column=1, value=int(t))


def build_result1():
    times = np.arange(0, 1800 + 1, 1)   # 0~1800, 步长1
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "温度"
    write_sheet(ws1, times, DIST)
    ws2 = wb.create_sheet("水分浓度")
    write_sheet(ws2, times, DIST)
    wb.save(DATA_DIR / "result1_template.xlsx")
    print("result1_template.xlsx 已生成，行数:", len(times) + 1)


def build_result2():
    # 时间范围待定，先按 3 小时 = 10800 s 生成
    times = np.arange(0, 10800 + 1, 1)
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "温度"
    write_sheet(ws1, times, DIST)
    ws2 = wb.create_sheet("水分浓度")
    write_sheet(ws2, times, DIST)
    wb.save(DATA_DIR / "result2_template.xlsx")
    print("result2_template.xlsx 已生成，行数:", len(times) + 1)


def build_result3():
    # 先按 72 h = 259200 s 生成，实际结束时间由甲给出
    times = np.arange(60, 259200 + 1, 60)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    write_sheet(ws, times, DIST)
    wb.save(DATA_DIR / "result3_template.xlsx")
    print("result3_template.xlsx 已生成，行数:", len(times) + 1)


def build_result4():
    times = np.arange(60, 259200 + 1, 60)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    write_sheet(ws, times, DIST, last_col_name="药材表面")
    wb.save(DATA_DIR / "result4_template.xlsx")
    print("result4_template.xlsx 已生成，行数:", len(times) + 1)


if __name__ == "__main__":
    build_result1()
    build_result2()
    build_result3()
    build_result4()
    print("\n全部模板已生成到:", DATA_DIR)